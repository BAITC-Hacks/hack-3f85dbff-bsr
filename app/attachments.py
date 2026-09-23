from __future__ import annotations

import base64
import uuid
from pathlib import Path

from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document

from .config import BASE_DIR

UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

# Для hackathon хранится в памяти. В production заменить на Redis/БД/Object Storage.
ATTACHMENTS: dict[str, dict] = {}


def _text_from_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)[:30000]


def _text_from_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)[:30000]


def _text_from_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    rows: list[str] = []
    for ws in wb.worksheets[:5]:
        rows.append(f"Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                rows.append(" | ".join(values))
            if sum(len(x) for x in rows) > 30000:
                break
    return "\n".join(rows)[:30000]


async def save_upload(file: UploadFile) -> dict:
    attachment_id = str(uuid.uuid4())
    suffix = Path(file.filename or "file").suffix.lower()
    path = UPLOAD_DIR / f"{attachment_id}{suffix}"
    content = await file.read()
    path.write_bytes(content)

    extracted_text = ""
    image_data_url = None
    if suffix == ".pdf":
        extracted_text = _text_from_pdf(path)
    elif suffix == ".docx":
        extracted_text = _text_from_docx(path)
    elif suffix in {".xlsx", ".xlsm"}:
        extracted_text = _text_from_xlsx(path)
    elif suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        mime = file.content_type or "image/jpeg"
        image_data_url = f"data:{mime};base64,{base64.b64encode(content).decode('ascii')}"
    else:
        extracted_text = content.decode("utf-8", errors="ignore")[:30000]

    item = {
        "id": attachment_id,
        "name": file.filename,
        "content_type": file.content_type,
        "text": extracted_text,
        "image_data_url": image_data_url,
    }
    ATTACHMENTS[attachment_id] = item
    return {"id": attachment_id, "name": file.filename, "has_text": bool(extracted_text), "is_image": bool(image_data_url)}


def get_attachments(ids: list[str]) -> list[dict]:
    return [ATTACHMENTS[x] for x in ids if x in ATTACHMENTS]
